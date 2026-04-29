"""
Analytics and Excel report generation module for Health Bot.
Creates weekly and monthly reports with charts and statistics.
"""

from typing import Dict, List, Any, Optional, Tuple
from datetime import date, datetime, timedelta
from pathlib import Path
import logging
import asyncio

from database.manager import db_manager
from database.models import (
    NutritionLog, TrainingLog, HealthMetric, User, DailyCalorieTarget,
    ExcelReport, Supplement
)
from config import config

logger = logging.getLogger(__name__)


class ExcelReportGenerator:
    """Generator for Excel reports with charts and analytics."""
    
    def __init__(self):
        self.db = db_manager
        self.reports_dir = config.REPORTS_DIR
        
        # Ensure reports directory exists
        Path(self.reports_dir).mkdir(parents=True, exist_ok=True)
    
    def generate_weekly_report(
        self,
        user_id: int,
        week_start: date = None
    ) -> Tuple[bool, str, str]:
        """
        Generate weekly Excel report.
        
        Args:
            user_id: User internal ID
            week_start: Start of week (default: most recent Monday)
        
        Returns:
            Tuple of (success, file_path, message)
        """
        try:
            import pandas as pd
            from openpyxl import Workbook
            from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
            from openpyxl.chart import LineChart, BarChart, Reference, PieChart
            from openpyxl.utils.dataframe import write_dataframe
            from openpyxl.utils import get_column_letter
        except ImportError:
            logger.error("pandas or openpyxl not installed")
            return False, "", "❌ Необходимые библиотеки не установлены"
        
        if week_start is None:
            today = date.today()
            week_start = today - timedelta(days=today.weekday())
        
        week_end = week_start + timedelta(days=6)
        
        try:
            # Create workbook
            wb = Workbook()
            wb.remove(wb.active)  # Remove default sheet
            
            # Get user
            user = self.db.get_user_by_id(user_id)
            if not user:
                return False, "", "❌ Пользователь не найден"
            
            # Generate sheets
            self._create_summary_sheet(wb, user, week_start, week_end)
            self._create_nutrition_sheet(wb, user_id, week_start, week_end)
            self._create_training_sheet(wb, user_id, week_start, week_end)
            self._create_health_sheet(wb, user_id, week_start, week_end)
            self._create_charts_sheet(wb, user_id, week_start, week_end)
            
            # Save file
            filename = f"weekly_report_{user_id}_{week_start.isoformat()}.xlsx"
            filepath = Path(self.reports_dir) / filename
            
            wb.save(filepath)
            
            # Log report
            self.db.add_excel_report(
                user_id=user_id,
                report_type="weekly",
                period_start=week_start,
                period_end=week_end,
                file_path=str(filepath)
            )
            
            return True, str(filepath), f"✅ Отчет за {week_start.strftime('%d.%m')} - {week_end.strftime('%d.%m.%Y')} готов"
            
        except Exception as e:
            logger.error(f"Error generating weekly report: {e}")
            import traceback
            traceback.print_exc()
            return False, "", f"❌ Ошибка генерации отчета: {e}"
    
    def generate_monthly_report(
        self,
        user_id: int,
        year: int = None,
        month: int = None
    ) -> Tuple[bool, str, str]:
        """
        Generate monthly Excel report.
        
        Args:
            user_id: User internal ID
            year: Year (default: current)
            month: Month (default: current)
        
        Returns:
            Tuple of (success, file_path, message)
        """
        try:
            import pandas as pd
            from openpyxl import Workbook
            from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
            from openpyxl.chart import LineChart, BarChart, Reference
        except ImportError:
            logger.error("pandas or openpyxl not installed")
            return False, "", "❌ Необходимые библиотеки не установлены"
        
        if year is None or month is None:
            today = date.today()
            year = today.year
            month = today.month
        
        # First and last day of month
        month_start = date(year, month, 1)
        if month == 12:
            month_end = date(year + 1, 1, 1) - timedelta(days=1)
        else:
            month_end = date(year, month + 1, 1) - timedelta(days=1)
        
        try:
            # Create workbook
            wb = Workbook()
            wb.remove(wb.active)
            
            # Get user
            user = self.db.get_user_by_id(user_id)
            if not user:
                return False, "", "❌ Пользователь не найден"
            
            # Generate sheets
            self._create_summary_sheet(wb, user, month_start, month_end, is_monthly=True)
            self._create_nutrition_sheet(wb, user_id, month_start, month_end, is_monthly=True)
            self._create_training_sheet(wb, user_id, month_start, month_end, is_monthly=True)
            self._create_health_sheet(wb, user_id, month_start, month_end, is_monthly=True)
            self._create_charts_sheet(wb, user_id, month_start, month_end, is_monthly=True)
            
            # Save file
            month_name = month_start.strftime('%B')
            filename = f"monthly_report_{user_id}_{year}_{month:02d}.xlsx"
            filepath = Path(self.reports_dir) / filename
            
            wb.save(filepath)
            
            # Log report
            self.db.add_excel_report(
                user_id=user_id,
                report_type="monthly",
                period_start=month_start,
                period_end=month_end,
                file_path=str(filepath)
            )
            
            return True, str(filepath), f"✅ Отчет за {month_name} {year} готов"
            
        except Exception as e:
            logger.error(f"Error generating monthly report: {e}")
            import traceback
            traceback.print_exc()
            return False, "", f"❌ Ошибка генерации отчета: {e}"
    
    def _create_summary_sheet(
        self,
        wb: Any,
        user: User,
        period_start: date,
        period_end: date,
        is_monthly: bool = False
    ):
        """Create summary sheet."""
        import pandas as pd
        from openpyxl.styles import Font, Alignment, PatternFill
        from openpyxl.utils.dataframe import write_dataframe
        
        ws = wb.create_sheet("📊 Сводка")
        
        # Title
        period_type = "Месячный" if is_monthly else "Недельный"
        ws['A1'] = f"{period_type} отчет о здоровье и тренировках"
        ws['A1'].font = Font(bold=True, size=16)
        
        # User info
        ws['A3'] = f"Пользователь: {user.name}"
        ws['A4'] = f"Период: {period_start.strftime('%d.%m.%Y')} - {period_end.strftime('%d.%m.%Y')}"
        ws['A5'] = f"Дата генерации: {datetime.now().strftime('%d.%m.%Y %H:%M')}"
        
        # Get statistics
        stats = self._calculate_period_statistics(user.id, period_start, period_end)
        
        # Summary data
        summary_data = [
            ["📈 ОБЩАЯ СТАТИСТИКА"],
            ["Дней в периоде", stats['days_in_period']],
            ["Дней с записями", stats['days_with_logs']],
            ["", ""],
            ["🍽️ ПИТАНИЕ (среднее за день)"],
            ["Калории", f"{stats['avg_calories']:.0f} ккал"],
            ["Белки", f"{stats['avg_protein']:.1f} г"],
            ["Жиры", f"{stats['avg_fat']:.1f} г"],
            ["Углеводы", f"{stats['avg_carbs']:.1f} г"],
            ["", ""],
            ["💪 ТРЕНИРОВКИ"],
            ["Всего тренировок", stats['total_workouts']],
            ["Общая длительность", f"{stats['total_training_minutes']} мин"],
            ["Сожжено калорий", f"{stats['total_calories_burned']:.0f} ккал"],
            ["", ""],
            ["❤️ ЗДОРОВЬЕ"],
            ["Вес (начало)", f"{stats['start_weight']} кг" if stats['start_weight'] else "Н/Д"],
            ["Вес (конец)", f"{stats['end_weight']} кг" if stats['end_weight'] else "Н/Д"],
            ["Изменение", f"{stats['weight_change']:+.1f} кг" if stats['weight_change'] else "Н/Д"],
        ]
        
        # Write data
        for row_idx, row_data in enumerate(summary_data, start=7):
            for col_idx, value in enumerate(row_data, start=1):
                cell = ws.cell(row=row_idx, column=col_idx, value=value)
                if col_idx == 1:
                    cell.font = Font(bold=True)
                cell.alignment = Alignment(horizontal='left', vertical='center')
        
        # Adjust column widths
        ws.column_dimensions['A'].width = 25
        ws.column_dimensions['B'].width = 20
    
    def _create_nutrition_sheet(
        self,
        wb: Any,
        user_id: int,
        period_start: date,
        period_end: date,
        is_monthly: bool = False
    ):
        """Create nutrition sheet."""
        import pandas as pd
        from openpyxl.utils.dataframe import write_dataframe
        from openpyxl.styles import PatternFill
        
        ws = wb.create_sheet("🍽️ Питание")
        
        # Get nutrition logs
        with self.db.session_scope() as session:
            logs = session.query(NutritionLog).filter(
                NutritionLog.user_id == user_id,
                NutritionLog.date >= period_start,
                NutritionLog.date <= period_end
            ).order_by(NutritionLog.date.asc()).all()
        
        if not logs:
            ws['A1'] = "Нет данных о питании за этот период"
            return
        
        # Aggregate by day
        daily_data = {}
        for log in logs:
            date_key = log.date.isoformat()
            if date_key not in daily_data:
                daily_data[date_key] = {
                    'date': log.date,
                    'calories': 0,
                    'protein': 0,
                    'fat': 0,
                    'carbs': 0,
                    'meals': 0
                }
            daily_data[date_key]['calories'] += log.total_calories or 0
            daily_data[date_key]['protein'] += log.total_protein or 0
            daily_data[date_key]['fat'] += log.total_fat or 0
            daily_data[date_key]['carbs'] += log.total_carbs or 0
            daily_data[date_key]['meals'] += 1
        
        # Create DataFrame
        df = pd.DataFrame(list(daily_data.values()))
        df['date'] = df['date'].dt.strftime('%d.%m.%Y')
        df = df[['date', 'calories', 'protein', 'fat', 'carbs', 'meals']]
        df.columns = ['Дата', 'Калории', 'Белки (г)', 'Жиры (г)', 'Углеводы (г)', 'Приемов пищи']
        
        # Write to sheet
        write_dataframe(ws, df, startcol=0, startrow=0, index=False)
        
        # Style header
        for cell in ws[1]:
            cell.fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
            cell.font = Font(bold=True, color="FFFFFF")
        
        # Adjust column widths
        for col in ws.columns:
            max_length = 0
            column = col[0].column_letter
            for cell in col:
                if cell.value:
                    max_length = max(max_length, len(str(cell.value)))
            ws.column_dimensions[column].width = min(max_length + 2, 20)
    
    def _create_training_sheet(
        self,
        wb: Any,
        user_id: int,
        period_start: date,
        period_end: date,
        is_monthly: bool = False
    ):
        """Create training sheet."""
        import pandas as pd
        from openpyxl.utils.dataframe import write_dataframe
        from openpyxl.styles import PatternFill
        
        ws = wb.create_sheet("💪 Тренировки")
        
        # Get training logs
        with self.db.session_scope() as session:
            logs = session.query(TrainingLog).filter(
                TrainingLog.user_id == user_id,
                TrainingLog.date >= period_start,
                TrainingLog.date <= period_end
            ).order_by(TrainingLog.date.asc()).all()
        
        if not logs:
            ws['A1'] = "Нет данных о тренировках за этот период"
            return
        
        # Prepare data
        data = []
        for log in logs:
            muscle_groups = log.muscle_groups or []
            groups_str = ", ".join([g.get('group', '') for g in muscle_groups])
            
            data.append({
                'date': log.date.strftime('%d.%m.%Y'),
                'duration': log.duration_minutes,
                'muscle_groups': groups_str,
                'calories': log.calories_burned or 0,
                'notes': log.notes[:50] if log.notes else ""
            })
        
        df = pd.DataFrame(data)
        df = df[['date', 'duration', 'muscle_groups', 'calories', 'notes']]
        df.columns = ['Дата', 'Длительность (мин)', 'Группы мышц', 'Калории', 'Заметки']
        
        # Write to sheet
        write_dataframe(ws, df, startcol=0, startrow=0, index=False)
        
        # Style header
        for cell in ws[1]:
            cell.fill = PatternFill(start_color="70AD47", end_color="70AD47", fill_type="solid")
            cell.font = Font(bold=True, color="FFFFFF")
        
        # Adjust column widths
        ws.column_dimensions['A'].width = 12
        ws.column_dimensions['B'].width = 15
        ws.column_dimensions['C'].width = 30
        ws.column_dimensions['D'].width = 12
        ws.column_dimensions['E'].width = 25
    
    def _create_health_sheet(
        self,
        wb: Any,
        user_id: int,
        period_start: date,
        period_end: date,
        is_monthly: bool = False
    ):
        """Create health metrics sheet."""
        import pandas as pd
        from openpyxl.utils.dataframe import write_dataframe
        from openpyxl.styles import PatternFill
        
        ws = wb.create_sheet("❤️ Здоровье")
        
        # Get health metrics
        with self.db.session_scope() as session:
            metrics = session.query(HealthMetric).filter(
                HealthMetric.user_id == user_id,
                HealthMetric.date >= period_start,
                HealthMetric.date <= period_end
            ).order_by(HealthMetric.date.asc()).all()
        
        if not metrics:
            ws['A1'] = "Нет данных о здоровье за этот период"
            return
        
        # Prepare data
        data = []
        for metric in metrics:
            data.append({
                'date': metric.date.strftime('%d.%m.%Y'),
                'weight': metric.weight_kg or '',
                'body_fat': metric.body_fat_percent or '',
                'heart_rate': metric.heart_rate_bpm or '',
                'sleep': (metric.sleep_duration_minutes or 0) // 60 if metric.sleep_duration_minutes else '',
                'steps': metric.steps or '',
                'stress': metric.stress_level or ''
            })
        
        df = pd.DataFrame(data)
        df = df[['date', 'weight', 'body_fat', 'heart_rate', 'sleep', 'steps', 'stress']]
        df.columns = ['Дата', 'Вес (кг)', 'Жир (%)', 'Пульс', 'Сон (ч)', 'Шаги', 'Стресс']
        
        # Write to sheet
        write_dataframe(ws, df, startcol=0, startrow=0, index=False)
        
        # Style header
        for cell in ws[1]:
            cell.fill = PatternFill(start_color="ED7D31", end_color="ED7D31", fill_type="solid")
            cell.font = Font(bold=True, color="FFFFFF")
        
        # Adjust column widths
        for col in ws.columns:
            max_length = 0
            column = col[0].column_letter
            for cell in col:
                if cell.value:
                    max_length = max(max_length, len(str(cell.value)))
            ws.column_dimensions[column].width = min(max_length + 2, 15)
    
    def _create_charts_sheet(
        self,
        wb: Any,
        user_id: int,
        period_start: date,
        period_end: date,
        is_monthly: bool = False
    ):
        """Create charts sheet."""
        from openpyxl.chart import LineChart, Reference, BarChart
        from openpyxl.styles import Font
        
        ws = wb.create_sheet("📈 Графики")
        
        # Get nutrition data for chart
        with self.db.session_scope() as session:
            nutrition_logs = session.query(NutritionLog).filter(
                NutritionLog.user_id == user_id,
                NutritionLog.date >= period_start,
                NutritionLog.date <= period_end
            ).order_by(NutritionLog.date.asc()).all()
        
        if nutrition_logs:
            # Aggregate by day
            daily_calories = {}
            for log in nutrition_logs:
                date_key = log.date.isoformat()
                if date_key not in daily_calories:
                    daily_calories[date_key] = 0
                daily_calories[date_key] += log.total_calories or 0
            
            # Create data for chart
            dates = sorted(daily_calories.keys())
            calories = [daily_calories[d] for d in dates]
            
            # Add data to sheet for chart reference
            ws['A1'] = "Дата"
            ws['B1'] = "Калории"
            for i, (d, c) in enumerate(zip(dates, calories), start=2):
                ws.cell(row=i, column=1, value=d)
                ws.cell(row=i, column=2, value=c)
            
            # Create line chart
            chart = LineChart()
            chart.title = "Динамика калорий по дням"
            chart.style = 10
            chart.y_axis.title = "Ккал"
            chart.x_axis.title = "Дата"
            
            data = Reference(ws, min_col=2, min_row=1, max_row=len(dates) + 1)
            cats = Reference(ws, min_col=1, min_row=2, max_row=len(dates) + 1)
            chart.add_data(data, titles_from_data=True)
            chart.set_categories(cats)
            
            ws.add_chart(chart, "D2")
        
        # Get training data for chart
        with self.db.session_scope() as session:
            training_logs = session.query(TrainingLog).filter(
                TrainingLog.user_id == user_id,
                TrainingLog.date >= period_start,
                TrainingLog.date <= period_end
            ).order_by(TrainingLog.date.asc()).all()
        
        if training_logs:
            # Create data for chart
            ws['A20'] = "Дата"
            ws['B20'] = "Длительность (мин)"
            for i, log in enumerate(training_logs, start=21):
                ws.cell(row=i, column=1, value=log.date.strftime('%d.%m'))
                ws.cell(row=i, column=2, value=log.duration_minutes)
            
            # Create bar chart
            chart = BarChart()
            chart.title = "Тренировки по дням"
            chart.style = 10
            chart.y_axis.title = "Минуты"
            chart.x_axis.title = "Дата"
            
            data = Reference(ws, min_col=2, min_row=19, max_row=len(training_logs) + 20)
            cats = Reference(ws, min_col=1, min_row=21, max_row=len(training_logs) + 20)
            chart.add_data(data, titles_from_data=True)
            chart.set_categories(cats)
            
            ws.add_chart(chart, "D20")
    
    def _calculate_period_statistics(
        self,
        user_id: int,
        period_start: date,
        period_end: date
    ) -> Dict[str, Any]:
        """Calculate statistics for the period."""
        stats = {
            'days_in_period': (period_end - period_start).days + 1,
            'days_with_logs': 0,
            'avg_calories': 0,
            'avg_protein': 0,
            'avg_fat': 0,
            'avg_carbs': 0,
            'total_workouts': 0,
            'total_training_minutes': 0,
            'total_calories_burned': 0,
            'start_weight': None,
            'end_weight': None,
            'weight_change': None
        }
        
        # Nutrition stats
        with self.db.session_scope() as session:
            daily_nutrition = session.query(
                NutritionLog.date,
                db_manager.session_factory.kw.get('bind').func.sum(NutritionLog.total_calories).label('calories'),
                db_manager.session_factory.kw.get('bind').func.sum(NutritionLog.total_protein).label('protein'),
                db_manager.session_factory.kw.get('bind').func.sum(NutritionLog.total_fat).label('fat'),
                db_manager.session_factory.kw.get('bind').func.sum(NutritionLog.total_carbs).label('carbs')
            ).filter(
                NutritionLog.user_id == user_id,
                NutritionLog.date >= period_start,
                NutritionLog.date <= period_end
            ).group_by(NutritionLog.date).all()
        
        if daily_nutrition:
            stats['days_with_logs'] = len(daily_nutrition)
            stats['avg_calories'] = sum(d.calories or 0 for d in daily_nutrition) / len(daily_nutrition)
            stats['avg_protein'] = sum(d.protein or 0 for d in daily_nutrition) / len(daily_nutrition)
            stats['avg_fat'] = sum(d.fat or 0 for d in daily_nutrition) / len(daily_nutrition)
            stats['avg_carbs'] = sum(d.carbs or 0 for d in daily_nutrition) / len(daily_nutrition)
        
        # Training stats
        with self.db.session_scope() as session:
            training_logs = session.query(TrainingLog).filter(
                TrainingLog.user_id == user_id,
                TrainingLog.date >= period_start,
                TrainingLog.date <= period_end
            ).all()
        
        if training_logs:
            stats['total_workouts'] = len(training_logs)
            stats['total_training_minutes'] = sum(log.duration_minutes or 0 for log in training_logs)
            stats['total_calories_burned'] = sum(log.calories_burned or 0 for log in training_logs)
        
        # Weight stats
        with self.db.session_scope() as session:
            weight_metrics = session.query(HealthMetric).filter(
                HealthMetric.user_id == user_id,
                HealthMetric.weight_kg != None,
                HealthMetric.date >= period_start,
                HealthMetric.date <= period_end
            ).order_by(HealthMetric.date.asc()).all()
        
        if weight_metrics:
            stats['start_weight'] = weight_metrics[0].weight_kg
            stats['end_weight'] = weight_metrics[-1].weight_kg
            stats['weight_change'] = stats['end_weight'] - stats['start_weight']
        
        return stats
    
    def get_user_reports(self, user_id: int) -> List[Dict[str, Any]]:
        """Get list of generated reports for user."""
        reports = self.db.get_excel_reports(user_id)
        
        return [
            {
                'id': r.id,
                'type': r.report_type,
                'period_start': r.period_start.isoformat(),
                'period_end': r.period_end.isoformat(),
                'file_path': r.file_path,
                'generated_at': r.generated_at.isoformat()
            }
            for r in reports
        ]


# Global instance
excel_report_generator = ExcelReportGenerator()


# ==================== ASYNC REPORT GENERATION ====================

async def generate_report_async(
    user_id: int,
    report_type: str = "weekly",
    callback=None
) -> Tuple[bool, str, str]:
    """
    Async wrapper for report generation.
    Runs report generation in background without blocking the bot.
    
    Args:
        user_id: User internal ID
        report_type: "weekly" or "monthly"
        callback: Optional callback function(success, file_path, message)
    
    Returns:
        Tuple of (success, file_path, message)
    """
    import asyncio
    
    def _generate():
        if report_type == "weekly":
            return excel_report_generator.generate_weekly_report(user_id)
        else:
            return excel_report_generator.generate_monthly_report(user_id)
    
    try:
        # Run in executor to not block
        loop = asyncio.get_event_loop()
        result = await loop.run_in_executor(None, _generate)
        
        if callback:
            await callback(*result)
        
        return result
        
    except Exception as e:
        logger.error(f"Error in async report generation: {e}")
        error_msg = (False, "", f"❌ Ошибка генерации отчета: {e}")
        if callback:
            await callback(*error_msg)
        return error_msg


class AsyncExcelReportGenerator:
    """Async Excel report generator with queue support."""
    
    def __init__(self):
        self.db = db_manager
        self.reports_dir = config.REPORTS_DIR
        self.generation_queue = asyncio.Queue()
        self.is_processing = False
        
        Path(self.reports_dir).mkdir(parents=True, exist_ok=True)
    
    async def queue_report(self, user_id: int, report_type: str = "weekly"):
        """Add report to generation queue."""
        await self.generation_queue.put({
            "user_id": user_id,
            "report_type": report_type,
            "created_at": datetime.utcnow()
        })
        
        # Start processing if not already running
        if not self.is_processing:
            asyncio.create_task(self.process_queue())
        
        return True, "", "🔄 Отчет добавлен в очередь на генерацию"
    
    async def process_queue(self):
        """Process report generation queue."""
        self.is_processing = True
        
        while not self.generation_queue.empty():
            try:
                task = await self.generation_queue.get()
                user_id = task["user_id"]
                report_type = task["report_type"]
                
                logger.info(f"Generating {report_type} report for user {user_id}")
                
                # Generate report in thread pool
                loop = asyncio.get_event_loop()
                result = await loop.run_in_executor(
                    None,
                    lambda: excel_report_generator.generate_weekly_report(user_id)
                    if report_type == "weekly"
                    else excel_report_generator.generate_monthly_report(user_id)
                )
                
                # Notify user (callback would be implemented here)
                logger.info(f"Report generated: {result}")
                
            except Exception as e:
                logger.error(f"Error processing report queue: {e}")
        
        self.is_processing = False
    
    def get_queue_status(self) -> Dict[str, Any]:
        """Get queue status."""
        return {
            "queue_size": self.generation_queue.qsize(),
            "is_processing": self.is_processing
        }


# Async instance
async_excel_generator = AsyncExcelReportGenerator()
report_generator = ExcelReportGenerator()
